"""
Unit tests for the Template Mass Update docx engine.

Runs with pytest *or* standalone (``py -3 test_docx_replace.py``) so it works
in this repo even though pytest isn't installed. Fixtures are built in-memory
with python-docx, plus a raw-XML text box injected to exercise the shape pass.

Covered:
  * run-split token in a body paragraph (join-runs fallback)
  * token inside a header paragraph
  * token split across <w:t> runs inside a text box (XML splice pass)
  * dry run counts matches but returns the original bytes unchanged
  * verification reports the document is clean after a real run
  * lookup CSV parsing (header row + positional)
"""

import io
import zipfile

from docx import Document

from backend.services.docx_replace import (
    Replacement,
    parse_lookup_csv,
    process_docx_bytes,
)
from backend.services.docx_verify import verify_docx_bytes


# ── fixture builders ────────────────────────────────────────────────────────

def _base_doc_bytes():
    """Body paragraph with a RUN-SPLIT token + a header with a token."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Hello << Ol")          # token deliberately split across two runs
    p.add_run("d Field >> world")
    doc.sections[0].header.paragraphs[0].add_run("Header has << Old Field >> too")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _inject_textbox(data: bytes, split_token=("Box &lt;&lt; Ol", "d Field &gt;&gt; end")):
    """Insert a well-formed <w:txbxContent> block (token split across <w:t>).

    The token parts are XML-escaped exactly like Word stores them (``<`` → ``&lt;``)
    so the injected document.xml stays well-formed.
    """
    src = io.BytesIO(data)
    with zipfile.ZipFile(src, "r") as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    xml = parts["word/document.xml"].decode("utf-8")
    txbx = (
        "<w:txbxContent><w:p>"
        f'<w:r><w:t xml:space="preserve">{split_token[0]}</w:t></w:r>'
        f'<w:r><w:t xml:space="preserve">{split_token[1]}</w:t></w:r>'
        "</w:p></w:txbxContent>"
    )
    xml = xml.replace("</w:body>", txbx + "</w:body>")
    parts["word/document.xml"] = xml.encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in parts.items():
            zout.writestr(name, content)
    return out.getvalue()


def _inject_hyperlink(data: bytes, text="e-service@quillarrowlaw.com"):
    """Insert a body paragraph whose text lives inside a <w:hyperlink> run.

    Reproduces the real-world case where an email/URL is a clickable mailto
    link: python-docx's ``paragraph.runs`` omits hyperlink-nested runs, so the
    engine used to report 0 matches on documents that clearly contained the text.
    """
    src = io.BytesIO(data)
    with zipfile.ZipFile(src, "r") as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    xml = parts["word/document.xml"].decode("utf-8")
    para = (
        "<w:p><w:hyperlink w:history=\"1\">"
        f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'
        "</w:hyperlink></w:p>"
    )
    xml = xml.replace("</w:body>", para + "</w:body>")
    parts["word/document.xml"] = xml.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in parts.items():
            zout.writestr(name, content)
    return out.getvalue()


RULES = [Replacement(old="<< Old Field >>", new="<< New Field >>")]


# ── tests ───────────────────────────────────────────────────────────────────

def test_run_split_body_and_header():
    data = _base_doc_bytes()
    res = process_docx_bytes(data, RULES)
    # one in body (run-split), one in header
    assert res.total == 2, res.per_rule
    assert res.changed
    doc = Document(io.BytesIO(res.data))
    assert "<< New Field >>" in doc.paragraphs[0].text
    assert "<< Old Field >>" not in doc.paragraphs[0].text
    hdr = doc.sections[0].header.paragraphs[0].text
    assert "<< New Field >>" in hdr


def test_textbox_xml_pass():
    data = _inject_textbox(_base_doc_bytes())
    res = process_docx_bytes(data, RULES)
    # body + header + textbox == 3
    assert res.total == 3, res.per_rule
    # raw xml should no longer contain the old token anywhere
    with zipfile.ZipFile(io.BytesIO(res.data)) as z:
        raw = z.read("word/document.xml").decode("utf-8")
    assert "Old Field" not in raw
    assert "New Field" in raw


def test_hyperlink_run_text_is_found_and_replaced():
    data = _inject_hyperlink(_base_doc_bytes())
    rules = [
        Replacement(
            old="e-service@quillarrowlaw.com",
            new="<< Matter.CustomField.EServiceEmail >>",
        )
    ]
    res = process_docx_bytes(data, rules)
    assert res.total == 1, res.per_rule
    with zipfile.ZipFile(io.BytesIO(res.data)) as z:
        raw = z.read("word/document.xml").decode("utf-8")
    assert "e-service@quillarrowlaw.com" not in raw
    assert "Matter.CustomField.EServiceEmail" in raw
    # the modified hyperlink is unwrapped (merge field left as plain text)
    assert "<w:hyperlink" not in raw
    # document still loads after the structural edit
    Document(io.BytesIO(res.data))


def test_dry_run_counts_but_does_not_change():
    data = _base_doc_bytes()
    res = process_docx_bytes(data, RULES, apply_changes=False)
    assert res.total == 2
    assert res.changed is False
    assert res.data == data  # original bytes untouched


def test_verify_clean_after_run():
    data = _inject_textbox(_base_doc_bytes())
    res = process_docx_bytes(data, RULES)
    report = verify_docx_bytes(res.data, RULES)
    assert report["clean"] is True, report
    assert report["old_remaining_total"] == 0
    assert report["new_present_total"] >= 3


def test_headers_footers_toggle_off():
    data = _base_doc_bytes()
    res = process_docx_bytes(data, RULES, include_headers_footers=False)
    assert res.total == 1  # only the body hit, header skipped


def test_parse_lookup_csv_header_and_positional():
    with_header = "old,new\r\nFoo,Bar\r\n<< A >>,<< B >>\r\n"
    rules = parse_lookup_csv(with_header)
    assert [(r.old, r.new) for r in rules] == [("Foo", "Bar"), ("<< A >>", "<< B >>")]

    positional = "Foo,Bar\r\nBaz,Qux\r\n"
    rules2 = parse_lookup_csv(positional)
    assert [(r.old, r.new) for r in rules2] == [("Foo", "Bar"), ("Baz", "Qux")]

    # blank old rows are skipped; blank value is allowed (clears nothing, just empties)
    messy = "old,new\r\n,ignored\r\nKeep,\r\n"
    rules3 = parse_lookup_csv(messy)
    assert [(r.old, r.new) for r in rules3] == [("Keep", "")]


def test_substring_old_inside_new_no_double_replace():
    """Regression: when old is a prefix/substring of new (e.g. Matter.Case.ID →
    Matter.Case.ID.TEST), the engine must NOT re-match the old text inside the
    new replacement, which would produce Matter.Case.ID.TEST.TEST."""
    doc = Document()
    doc.add_paragraph("Matter.Case.ID")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    rules = [Replacement(old="Matter.Case.ID", new="Matter.Case.ID.TEST")]
    res = process_docx_bytes(data, rules)
    result_doc = Document(io.BytesIO(res.data))
    result_text = result_doc.paragraphs[0].text
    assert result_text == "Matter.Case.ID.TEST", f"Got double replacement: {result_text!r}"
    assert res.total == 1, f"Expected 1 replacement, got {res.total}"


def test_partial_match_merge_fields_not_affected():
    """When replacing a merge field name that is a prefix of another field,
    the longer field must NOT be corrupted.

    E.g. replacing ``Matter.ResponsibleAttorney`` must NOT touch
    ``Matter.ResponsibleAttorney.JobTitle.Name`` or
    ``Matter.ResponsibleAttorney.Email``.
    """
    doc = Document()
    doc.add_paragraph("<< Matter.ResponsibleAttorney >>")
    doc.add_paragraph("<< Matter.ResponsibleAttorney.JobTitle.Name >>")
    doc.add_paragraph("<< Matter.ResponsibleAttorney.Email >>")
    doc.add_paragraph("<<Matter.ResponsibleAttorney>>")  # no spaces
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    rules = [Replacement(old="Matter.ResponsibleAttorney", new="Matter.ResponsibleAttorney.TEST")]
    res = process_docx_bytes(data, rules)
    result_doc = Document(io.BytesIO(res.data))
    texts = [p.text for p in result_doc.paragraphs if p.text.strip()]
    # Only the exact field should be replaced, not the longer fields
    assert "<< Matter.ResponsibleAttorney.TEST >>" == texts[0] or "Matter.ResponsibleAttorney.TEST" in texts[0], f"Expected replacement: {texts[0]!r}"
    assert "Matter.ResponsibleAttorney.JobTitle.Name" in texts[1], f"Longer field corrupted: {texts[1]!r}"
    assert "Matter.ResponsibleAttorney.Email" in texts[2], f"Email field corrupted: {texts[2]!r}"
    assert "Matter.ResponsibleAttorney.TEST" in texts[3], f"No-space variant not replaced: {texts[3]!r}"


def _run_all():
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    passed = 0
    for fn in fns:
        fn()
        print(f"  PASS {fn.__name__}")
        passed += 1
    print(f"\n{passed}/{len(fns)} tests passed.")


if __name__ == "__main__":
    _run_all()
